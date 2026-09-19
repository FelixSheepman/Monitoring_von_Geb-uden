"""Word-Berichtsexport (.docx): Datenpruefung, Abbildungen mit Beschriftung, Auswertung, Vergleich."""

from __future__ import annotations

import io
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN, RED, HEADER = "C6EFCE", "FFC7CE", "1F3864"


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _table(doc: Document, df: pd.DataFrame, col_widths_cm: list[float], status_col: str | None = None,
           font_pt: int = 8) -> None:
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    for j, name in enumerate(df.columns):
        c = t.rows[0].cells[j]
        c.text = ""
        run = c.paragraphs[0].add_run(str(name))
        run.bold, run.font.size, run.font.color.rgb = True, Pt(font_pt), RGBColor(255, 255, 255)
        _shade(c, HEADER)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for j, name in enumerate(df.columns):
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run("" if pd.isna(row[name]) else str(row[name]))
            run.font.size = Pt(font_pt)
            if status_col and name == status_col:
                good = str(row[name]).startswith(("Plausibel", "übereinstimmend"))
                _shade(cells[j], GREEN if good else RED)
    for row in t.rows:
        for j, w in enumerate(col_widths_cm):
            row.cells[j].width = Cm(w)


def _figure_png(fig) -> bytes | None:
    try:
        return fig.to_image(format="png", width=1000, height=480, scale=2)
    except Exception:
        return None


def export_docx(report, path_or_buffer, narrative=None, comparison=None, title="Monitoringbericht",
                subtitle="Automatisierte Auswertung durch den KI-Agenten") -> list[str]:
    """Schreibt den Bericht. `narrative`: Liste von NarrativeBlock (Default: regelbasiert aus report).
    `comparison`: optional (Tabelle1-Vergleich als DataFrame). Gibt Warnungen zurueck (z.B. fehlende Bilder)."""
    warnings: list[str] = []
    narrative = narrative if narrative is not None else report.narrative
    df = report.df

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(title, 0)
    doc.add_paragraph(subtitle)
    doc.add_paragraph(
        f"Messzeitraum: {df.index.min():%d.%m.%Y} bis {df.index.max():%d.%m.%Y} · {len(df):,} Zeitschritte (15 min) · "
        f"{df.shape[1]} Messspalten · erstellt am {date.today():%d.%m.%Y}".replace(",", ".")
    )

    doc.add_heading("1 Datenprüfung", 1)
    doc.add_paragraph(
        "Die Datenprüfung wurde automatisiert auf den Rohdaten (15-Minuten-Auflösung) durchgeführt. "
        "Grün markierte Spalten sind plausibel, rot markierte weisen Auffälligkeiten auf."
    )
    q = report.quality_df[["Spalte", "Bezeichnung", "Einheit", "Min", "Max", "Plausibilität", "Bewertung"]]
    cap = doc.add_paragraph()
    cap.add_run("Tabelle 1: Datenprüfung").bold = True
    _table(doc, q, [1.0, 4.6, 1.2, 1.5, 1.5, 1.8, 5.0], status_col="Plausibilität", font_pt=7)

    doc.add_heading("2 Grafische Aufbereitung", 1)
    by_key: dict[str, list] = {}
    for b in narrative:
        for k in b.figure_keys:
            by_key.setdefault(k, []).append(b)
    missing_images = 0
    for i, entry in enumerate(report.figures, start=1):
        png = _figure_png(entry.figure)
        if png:
            doc.add_picture(io.BytesIO(png), width=Cm(16))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            missing_images += 1
            doc.add_paragraph("[Abbildung konnte nicht gerendert werden]")
        c = doc.add_paragraph()
        c.add_run(f"Abbildung {i}: {entry.title}. ").bold = True
        c.add_run(entry.caption)
    if missing_images:
        warnings.append(
            f"{missing_images} Abbildung(en) konnten nicht als Bild eingefügt werden "
            "(Diagramm-Rendering benötigt einen installierten Chrome/Chromium-Browser)."
        )

    doc.add_heading("3 Auswertung der Ergebnisse", 1)
    for b in narrative:
        doc.add_heading(b.heading, 2)
        doc.add_paragraph(b.text)
        refs = [f"Abbildung {i}" for i, e in enumerate(report.figures, start=1) if e.key in b.figure_keys]
        if refs:
            p = doc.add_paragraph()
            r = p.add_run("Bezug: " + ", ".join(refs))
            r.italic, r.font.size = True, Pt(9)

    if comparison is not None and len(comparison):
        doc.add_heading("4 Vergleich manuelle Auswertung / Agent", 1)
        cap = doc.add_paragraph()
        cap.add_run("Tabelle 2: Gegenüberstellung der Datenprüfung").bold = True
        cmp_df = comparison[["Spalte", "Manuell", "Agent", "Ergebnis", "Manueller_Befund", "Agent_Befund"]]
        _table(doc, cmp_df, [1.0, 1.8, 1.8, 2.6, 4.2, 4.2], status_col="Ergebnis", font_pt=7)

    doc.save(path_or_buffer)
    return warnings

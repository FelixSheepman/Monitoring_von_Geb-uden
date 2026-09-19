"""Tests: Der Agent muss die Kennwerte der manuellen Tabelle 1 reproduzieren."""

from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from monitoring_agent.comparison import agreement_summary, compare_findings, compare_table1
from monitoring_agent.config import COLUMNS
from monitoring_agent.data_loader import infer_interval_minutes, load_measurements
from monitoring_agent.quality import quality_to_dataframe, run_data_quality
from monitoring_agent.report import build_report

DATA = Path(__file__).resolve().parent.parent / "source_docs" / "Messdaten 2024-2026.xlsx"
needs_data = pytest.mark.skipif(not DATA.exists(), reason="Messdaten-Datei fehlt")


@pytest.fixture(scope="module")
def df():
    return load_measurements(DATA)


@pytest.fixture(scope="module")
def report(df):
    return build_report(df)


@needs_data
def test_dataset_shape_matches_hausarbeit(df):
    assert df.shape == (58125, 23)
    assert df.index.min() == pd.Timestamp("2024-11-01 00:00:00")
    assert infer_interval_minutes(df) == 15


@needs_data
def test_min_max_match_manual_header_rows(df):
    wb = openpyxl.load_workbook(DATA, read_only=True, data_only=True)
    rows = list(wb["Tabelle1"].iter_rows(min_row=4, max_row=5, values_only=True))
    manual_min, manual_max = rows[0][1:24], rows[1][1:24]
    for i, col in enumerate(COLUMNS):
        assert df[col.short].min() == pytest.approx(manual_min[i], abs=0.01), col.short
        assert df[col.short].max() == pytest.approx(manual_max[i], abs=0.01), col.short


@needs_data
def test_dropout_columns_flagged_like_manual_table(report):
    q = report.quality_df.set_index("Spalte")
    for spalte in (12, 13, 14, 15, 18):
        assert q.loc[spalte, "Plausibilität"] == "Auffällig!"
        assert "Aussetzer" in q.loc[spalte, "Bewertung"]


@needs_data
def test_binary_pump_columns_plausible(report):
    q = report.quality_df.set_index("Spalte")
    assert q.loc[16, "Plausibilität"] == "Plausibel"
    assert q.loc[17, "Plausibilität"] == "Plausibel"


@needs_data
def test_all_manual_findings_reproduced(report):
    found = compare_findings([b.heading for b in report.narrative])
    assert (found["Vom Agent gefunden"] == "ja").all()


@needs_data
def test_table1_agreement_is_reported(report):
    summary = agreement_summary(compare_table1(report.quality_df))
    assert summary["Spalten gesamt"] == 23
    assert summary["übereinstimmend"] >= 13


@needs_data
def test_display_resampling_does_not_change_quality(df):
    raw = build_report(df).quality_df
    smoothed = build_report(df, display_resample="D").quality_df
    pd.testing.assert_frame_equal(raw, smoothed)


def _synthetic(**cols):
    idx = pd.date_range("2025-01-01", periods=200, freq="15min")
    base = {c.short: np.full(200, 30.0) for c in COLUMNS}
    for c in COLUMNS:
        if c.role == "pump":
            base[c.short] = np.ones(200)
        if c.role.startswith("meter"):
            base[c.short] = np.arange(200, dtype=float)
    base.update(cols)
    return pd.DataFrame(base, index=idx)


def test_single_zero_in_flow_temperature_is_flagged():
    vl = np.full(200, 40.0)
    vl[50] = 0
    q = quality_to_dataframe(run_data_quality(_synthetic(**{"RLT primär VL": vl}))).set_index("Spalte")
    assert q.loc[2, "Plausibilität"] == "Auffällig!"


def test_meter_reset_is_flagged():
    meter = np.arange(200, dtype=float)
    meter[100:] -= 50
    q = quality_to_dataframe(run_data_quality(_synthetic(**{"Zähler 019 – WMZ": meter}))).set_index("Spalte")
    assert q.loc[1, "Plausibilität"] == "Auffällig!"


def test_clean_synthetic_data_has_no_meter_or_zero_findings():
    q = quality_to_dataframe(run_data_quality(_synthetic()))
    assert not q["Bewertung"].str.contains("Aussetzer|rückläufig").any()


# --- LLM-Agent (ohne echten API-Aufruf) ---

from types import SimpleNamespace

from monitoring_agent.llm_agent import build_facts, generate_narrative, parse_blocks


class _FakeClient:
    def __init__(self, text, stop_reason="end_turn"):
        self.calls = []
        self._text, self._stop = text, stop_reason
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(
            stop_reason=self._stop, content=[SimpleNamespace(type="text", text=self._text)],
            usage=SimpleNamespace(input_tokens=1000, output_tokens=500))


def test_parse_blocks_filters_unknown_figure_keys():
    blocks = parse_blocks('Hier: [{"heading": "A", "figure_keys": ["heizkurve", "gibtsnicht"], "text": "T"}]')
    assert blocks[0].figure_keys == ["heizkurve"] and blocks[0].heading == "A"


@needs_data
def test_llm_agent_uses_only_computed_facts_and_fallback_for_opus(report):
    facts = build_facts(report)
    assert "Heizkurve ohne erkennbare Heizgrenze" in facts and "96 %" in facts
    fake = _FakeClient('[{"heading": "X", "figure_keys": [], "text": "Y"}]')
    result = generate_narrative(facts, fake, "claude-opus-5")
    assert result.blocks[0].heading == "X" and result.output_tokens == 500
    assert fake.calls[0]["extra_body"] == {"fallbacks": "default"}
    assert "Fakten" in fake.calls[0]["messages"][0]["content"]


def test_llm_agent_refusal_raises():
    with pytest.raises(RuntimeError):
        generate_narrative("{}", _FakeClient("[]", stop_reason="refusal"), "claude-sonnet-5")


# --- Erweiterungen: Regelsatz v2, Extras, Schwellenwerte ---

from monitoring_agent.extras import availability_daily, pump_runtime_monthly, savings_potential
from monitoring_agent.settings import Thresholds


@needs_data
def test_strict_rules_find_all_manual_findings(df):
    v1 = agreement_summary(compare_table1(quality_to_dataframe(run_data_quality(df))))
    v2 = agreement_summary(compare_table1(quality_to_dataframe(run_data_quality(df, strict=True))))
    key = "Trefferquote % (manuelle Auffälligkeiten vom Agent gefunden)"
    assert v2[key] == 100.0 and v2[key] > v1[key]
    assert v2["nur manuell auffällig"] == 0


@needs_data
def test_pumps_run_in_winter_and_mostly_off_in_summer(df):
    monthly = pump_runtime_monthly(df)
    share = monthly.div(monthly.index.days_in_month * 24, axis=0)
    winter = share[share.index.month.isin([12, 1, 2])]
    summer = share[share.index.month.isin([7])]
    assert (winter > 0.95).all().all()
    assert (summer < 0.5).all().all()


@needs_data
def test_pump_narrative_matches_data_not_a_fixed_claim(report):
    block = [b for b in report.narrative if b.figure_keys == ["pumpenlaufzeit"]][0]
    assert "Sommerabschaltung" in block.heading and "ganzjährig" not in block.heading


@needs_data
def test_savings_scale_with_assumptions(df):
    base = savings_potential(df, Thresholds())
    doubled = savings_potential(df, Thresholds(rlt_night_hours=16.0))
    assert doubled.loc[1, "Einsparung kWh/a"] == pytest.approx(2 * base.loc[1, "Einsparung kWh/a"], rel=0.01)
    assert base.loc[base["Maßnahme"] == "Summe", "Anteil %"].iloc[0] < 10


@needs_data
def test_availability_counts_known_zero_dropouts(df):
    assert availability_daily(df)["FBH Geb.06 VL (Ist)"].sum() >= 3


@needs_data
def test_delta_t_threshold_changes_narrative(df):
    strict = [b for b in build_report(df, thresholds=Thresholds(delta_t_min=5.0)).narrative if "spreizung" in b.heading][0]
    assert "über 5 Kelvin" in strict.text


@needs_data
def test_word_report_places_text_directly_below_its_figure(df, monkeypatch, tmp_path):
    from docx import Document
    from monitoring_agent import word_export
    monkeypatch.setattr(word_export, "_figure_png", lambda fig: None)
    rep = build_report(df)
    path = tmp_path / "b.docx"
    word_export.export_docx(rep, str(path))
    paras = [p.text for p in Document(str(path)).paragraphs]
    i = next(k for k, t in enumerate(paras) if t.startswith("Abbildung 3:"))
    assert paras[i + 1] == "Heizkurve ohne erkennbare Heizgrenze"
    assert not any(t == "3 Auswertung der Ergebnisse" for t in paras)


@needs_data
def test_word_report_explains_savings_measures(df, monkeypatch, tmp_path):
    from docx import Document
    from monitoring_agent import word_export
    monkeypatch.setattr(word_export, "_figure_png", lambda fig: None)
    path = tmp_path / "b.docx"
    word_export.export_docx(build_report(df), str(path))
    text = " ".join(p.text for p in Document(str(path)).paragraphs)
    for needle in ("Maßnahme 1: Heizgrenze", "Maßnahme 2: Nachtabsenkung", "Gesamteinordnung", "dritten Potenz"):
        assert needle in text

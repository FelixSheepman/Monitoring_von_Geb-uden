"""Tests fuer Kontrollkriterien, Bewertung, Datenabdeckung, Theorie-Praxis-Abgleich und Kapitel-8-Export."""

from pathlib import Path

import pandas as pd
import pytest
from docx import Document

from monitoring_agent import word_export
from monitoring_agent.assessment import (assessment_texts, build_assessment, data_coverage, measurement_head,
                                         sensor_overview, theory_vs_practice)
from monitoring_agent.config import COLUMNS
from monitoring_agent.data_loader import load_measurements, read_manual_minmax
from monitoring_agent.process import (STATUS_FAIL, STATUS_NA, STATUS_OK, Ctx, evaluate_criteria, load_criteria,
                                      optimization_hints, overall_summary, research_answers, step_report, step_summary)
from monitoring_agent.report import build_report
from monitoring_agent.structure import REPORT_SECTIONS, REQUIRED_SECTION_KEYWORDS, STEPS

DATA = Path(__file__).resolve().parent.parent / "source_docs" / "Messdaten 2024-2026.xlsx"
needs_data = pytest.mark.skipif(not DATA.exists(), reason="Messdaten-Datei fehlt")


@pytest.fixture(scope="module")
def df():
    return load_measurements(DATA)


@pytest.fixture(scope="module")
def report(df):
    return build_report(df)


@pytest.fixture(scope="module")
def ctx(report):
    return Ctx(report=report, manual_minmax=read_manual_minmax(DATA), assessment=report.assessment, coverage=report.coverage)


@pytest.fixture(scope="module")
def crit(ctx):
    return evaluate_criteria(ctx)


# ----------------------------------------------------------------------------- Kontrollkriterien

@needs_data
def test_every_step_has_criteria_and_all_are_evaluated(crit):
    assert set(crit["Schritt"]) == set(STEPS)
    assert len(crit) == len(load_criteria())


@needs_data
def test_known_criteria_results(crit):
    c = crit.set_index("ID")
    assert c.loc["DP2", "Ist"] == 100 and c.loc["DP2", "Status"] == STATUS_OK   # 46 von 46 Min/Max identisch
    assert c.loc["DP3", "Ist"] == 70 and c.loc["DP3", "Status"] == STATUS_FAIL  # Trefferquote v1
    assert c.loc["GR1", "Ist"] == 100                                           # alle 12 manuellen Abbildungen
    assert c.loc["GR3", "Status"] == STATUS_OK                                  # Kap. 5.2 eingehalten
    assert c.loc["BW4", "Status"] == STATUS_NA                                  # keine manuelle Bewertung hinterlegt


@needs_data
def test_reference_vs_self_check_is_labelled(crit):
    art = crit.set_index("ID")["Art"]
    assert art["DP3"] == "Referenzvergleich" and art["BW1"] == "Eigenprüfung"
    summ, text = overall_summary(crit)
    assert "Eigenprüfungen" in text
    assert summ.set_index("Zwischenschritt").loc["Bewertung", "davon Referenzvergleich"] == 0


@needs_data
def test_threshold_from_csv_changes_status(ctx):
    crit = load_criteria()
    crit.loc[crit["ID"] == "DP3", "Grenzwert"] = 60
    result = evaluate_criteria(ctx, crit).set_index("ID")
    assert result.loc["DP3", "Status"] == STATUS_OK


@needs_data
def test_optimization_hint_quantifies_rule_set_v2(crit, ctx):
    hints = optimization_hints(crit, ctx)
    assert "100 %" in hints["DP3"] and "BW4" in hints


@needs_data
def test_step_report_and_summary_are_consistent(crit):
    rep = step_report("Datenprüfung", crit, {})
    assert "Abweichungen" in rep["vergleich"] and len(rep["gegenueberstellung"]) == 6
    assert step_summary(crit)["Kriterien"].sum() == len(crit)


@needs_data
def test_research_answers_use_real_numbers(ctx, crit, report):
    from monitoring_agent.comparison import agreement_summary, compare_table1
    answers = research_answers(ctx, crit, agreement_summary(compare_table1(report.quality_df)), report.timings)
    assert set(answers) == {"FF1", "FF2", "FF3", "FF4"}
    assert "7.1 %" in answers["FF3"] and "6.854 kWh pro Jahr (1.923 € pro Jahr), das sind" in answers["FF3"]
    assert "manuelle Zeitaufwand" in answers["FF2"]


# ----------------------------------------------------------------------------- Bewertung, Abdeckung, Theorie

@needs_data
def test_assessment_is_complete_and_transparent(report):
    a = report.assessment
    assert (a["Einstufungsregel"] != "").all() and (a["Empfehlung"] != "").all()
    assert set(a["Schweregrad"]) <= {"hoch", "mittel", "gering"}
    by_key = a.set_index("Schlüssel")["Schweregrad"]
    assert by_key["heizkurve"] == "hoch" and by_key["rlt_dauerbetrieb"] == "hoch"
    assert by_key["datenqualitaet"] == "gering" and by_key["zaehler"] == "mittel"
    order = a["Schweregrad"].map({"hoch": 0, "mittel": 1, "gering": 2}).tolist()
    assert order == sorted(order)


@needs_data
def test_assessment_text_mentions_gap_to_target(report):
    text = assessment_texts(report.assessment, report.savings)["zusammenfassung"]
    assert "10 %" in text and "verfehlt" in text


@needs_data
def test_data_coverage_matches_dataset(df):
    cov = data_coverage(df)
    heiz = cov[cov["Zeitraum"] == "Heizperiode"]
    assert list(heiz["Erfüllt"]) == ["ja", "ja"]
    assert heiz.iloc[0]["Abdeckung %"] == pytest.approx(85.4, abs=0.1)   # Oktober 2024 fehlt
    assert (cov[cov["Zeitraum"] == "Sommerperiode"]["Erfüllt"] == "ja").sum() == 1


@needs_data
def test_theory_vs_practice_verdicts(df, report):
    tp = theory_vs_practice(df, report.thresholds, report.savings).set_index("Quelle", append=False)
    verdicts = " | ".join(tp["Ergebnis"])
    assert "nicht bestätigt" in verdicts and "teilweise bestätigt" in verdicts and "bestätigt" in verdicts
    assert len(tp) == 7


@needs_data
def test_measurement_head_and_sensor_overview(df):
    head = measurement_head(df, 5)
    assert len(head) == 5 and head.iloc[0]["Datum"] == "01.11.2024 00:00"
    assert len(sensor_overview()) == len(COLUMNS)


def test_manual_minmax_missing_returns_none(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabelle1"
    ws.append(["a"])
    path = tmp_path / "x.xlsx"
    wb.save(path)
    assert read_manual_minmax(path) is None


# ----------------------------------------------------------------------------- Word

@needs_data
def test_report_follows_required_structure_and_thesis_numbering(df, monkeypatch, tmp_path):
    monkeypatch.setattr(word_export, "_figure_png", lambda fig: None)
    path = tmp_path / "b.docx"
    word_export.export_docx(build_report(df), str(path))
    paras = [p.text for p in Document(str(path)).paragraphs]
    headings = [p.text for p in Document(str(path)).paragraphs if p.style.name == "Heading 1"]
    for keyword in REQUIRED_SECTION_KEYWORDS:
        assert any(keyword in h for h in headings), keyword
    assert [h.split(" ", 1)[1] for h in headings[:len(REPORT_SECTIONS)]] == REPORT_SECTIONS
    assert any(t.startswith("Abbildung 1: Messdatenkopf") for t in paras)
    assert any(t.startswith("Abbildung 4: Heizkurve") for t in paras)
    assert any(t.startswith("Abbildung 5: RLT primär VL") for t in paras)


@needs_data
def test_chapter8_document_has_thesis_structure(crit, ctx, report, tmp_path):
    path = tmp_path / "k8.docx"
    notes = {"Datenprüfung": "Meine Diskussion zur Datenprüfung"}
    word_export.export_chapter8_docx(str(path), crit, optimization_hints(crit, ctx), notes, report.timings)
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    for i in range(2, 2 + len(STEPS)):
        for suffix in ("1 Gegenüberstellung", "2 Vergleich", "3 Optimierung", "4 Diskussion", "5 Zusammenfassung"):
            assert f"8.{i}.{suffix}" in text
    assert "Meine Diskussion zur Datenprüfung" in text
    assert "[Diskussion von den Bearbeitern zu ergänzen]" in text      # Schritte ohne eigene Eingabe
    assert "Werkzeugvergleich: Excel und Python" in text
    assert "8.7 Zusammenfassung über alle Schritte" in text


@needs_data
def test_criteria_do_not_depend_on_anomaly_markers(df):
    """Die roten Anomalie-Marker duerfen die Bewertung der Grafikerstellung nicht veraendern."""
    with_markers = build_report(df, show_anomalies=True)
    c = evaluate_criteria(Ctx(report=with_markers, manual_minmax=read_manual_minmax(DATA),
                              assessment=with_markers.assessment)).set_index("ID")
    for id_ in ("GR1", "GR2", "GR3", "GR4"):
        assert c.loc[id_, "Status"] == STATUS_OK, (id_, c.loc[id_, "Detail"])

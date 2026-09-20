"""Tests fuer den Ausschluss fehlerhafter Einzelwerte (Auswahl, Wirkung auf Grafiken/Bewertung, Protokoll, Export)."""

from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from docx import Document

from monitoring_agent import word_export
from monitoring_agent.data_loader import load_measurements
from monitoring_agent.exclusion import RULES, all_of, apply_exclusions, find_invalid
from monitoring_agent.report import build_report

DATA = Path(__file__).resolve().parent.parent / "source_docs" / "Messdaten 2024-2026.xlsx"
needs_data = pytest.mark.skipif(not DATA.exists(), reason="Messdaten-Datei fehlt")
SOLL = "Stat. Heizung Geb.06 VL (Soll)"
WMZ = "Zähler 019 – WMZ"


@pytest.fixture(scope="module")
def df():
    return load_measurements(DATA)


@pytest.fixture(scope="module")
def cands(df):
    return find_invalid(df)


@needs_data
def test_candidates_list_known_faults_with_reason(cands):
    zero = cands[(cands["Spalte"] == SOLL) & (cands["rule_key"] == "nullwert")].iloc[0]
    assert zero["Anzahl"] == 16 and "Aussetzer" in zero["Begründung"]
    rew = cands[(cands["Spalte"] == WMZ) & (cands["rule_key"] == "zaehler_rueckgang")].iloc[0]
    assert rew["Anzahl"] == 3
    assert set(cands["rule_key"]) <= set(RULES) and (cands["Begründung"] != "").all()


@needs_data
def test_no_selection_keeps_everything_as_before(df):
    clean, log = apply_exclusions(df, frozenset())
    assert clean is df and log.n_values == 0 and log.summary.empty
    rep = build_report(df)
    assert rep.exclusion_log.n_values == 0 and rep.df.equals(rep.raw_df)


@needs_data
def test_exclusion_removes_values_from_analysis_but_not_from_quality_check(df, cands):
    sel = frozenset({(SOLL, "nullwert")})
    clean, log = apply_exclusions(df, sel)
    assert (df[SOLL] == 0).sum() == 16 and (clean[SOLL] == 0).sum() == 0     # Rohdaten bleiben unveraendert
    assert clean[SOLL].isna().sum() == df[SOLL].isna().sum() + 16
    assert log.n_values == 16 and log.n_columns == 1 and len(log.details) == 16
    assert (log.details["Messwert"] == 0).all() and log.details["Begründung"].str.len().min() > 20

    plain, excl = build_report(df), build_report(df, excluded=sel)
    assert excl.quality_df.equals(plain.quality_df)                          # Datenpruefung zeigt weiter die Rohdaten
    assert excl.raw_df[SOLL].min() == 0 and excl.df[SOLL].min() > 0          # Auswertung rechnet ohne die Nullwerte
    assert len(excl.figures) == len(plain.figures)


@needs_data
def test_excluded_values_do_not_appear_in_figures(df):
    excl = build_report(df, excluded=frozenset({(SOLL, "nullwert")}))
    fig = {e.key: e.figure for e in excl.figures}["regelguete_stat_heizung"]
    soll_trace = next(t for t in fig.data if "Soll" in (t.name or ""))
    assert np.nanmin(np.asarray(soll_trace.y, dtype=float)) > 0


@needs_data
def test_all_selected_and_overlap_is_counted_once(df, cands):
    clean, log = apply_exclusions(df, all_of(cands))
    assert log.n_values == int(log.details.shape[0]) == int(log.summary["Anzahl"].sum())
    assert not log.details.duplicated(["Zeitpunkt", "Spalte"]).any()
    assert clean.notna().sum().sum() == df.notna().sum().sum() - log.n_values
    assert set(log.summary["Regel"]) <= set(RULES.values())


@needs_data
def test_exclusion_is_part_of_excel_and_word_export(df, tmp_path, monkeypatch):
    import openpyxl
    from monitoring_agent.excel_export import export_workbook
    rep = build_report(df, excluded=frozenset({(SOLL, "nullwert")}))
    xlsx = tmp_path / "r.xlsx"
    export_workbook(rep, str(xlsx))
    assert "Ausgeschlossene Werte" in openpyxl.load_workbook(xlsx, read_only=True).sheetnames

    monkeypatch.setattr(word_export, "_figure_png", lambda fig: None)
    path = tmp_path / "b.docx"
    word_export.export_docx(rep, str(path))
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Nicht berücksichtigte Messwerte" in text and "16 fehlerhafte Messwerte" in text
    captions = [p.text for p in doc.paragraphs if p.text.startswith("Tabelle ")]
    assert [c.split(":")[0] for c in captions] == [f"Tabelle {i}" for i in range(1, len(captions) + 1)]

    plain = tmp_path / "p.docx"
    word_export.export_docx(build_report(df), str(plain))
    assert "Nicht berücksichtigte Messwerte" not in "\n".join(p.text for p in Document(str(plain)).paragraphs)


def test_synthetic_range_and_pair_rules():
    from monitoring_agent.config import COLUMNS
    idx = pd.date_range("2025-01-01", periods=6, freq="15min")
    frame = pd.DataFrame({c.short: np.linspace(30, 40, 6) for c in COLUMNS}, index=idx)
    frame["RLT KL01 Außenluft"] = [5, 5, 99, 5, 5, 5]                        # ueber Grenze 45 °C
    vl = next(c for c in COLUMNS if c.role == "vl")
    frame.loc[idx[1], vl.short] = 0                                          # Nullwert
    found = find_invalid(frame)
    assert ("RLT KL01 Außenluft", "grenzwert") in set(zip(found["Spalte"], found["rule_key"]))
    assert (vl.short, "nullwert") in set(zip(found["Spalte"], found["rule_key"]))
    clean, log = apply_exclusions(frame, all_of(found))
    assert np.isnan(clean.loc[idx[2], "RLT KL01 Außenluft"]) and np.isnan(clean.loc[idx[1], vl.short])
    assert frame.loc[idx[2], "RLT KL01 Außenluft"] == 99                     # Original bleibt erhalten


@needs_data
def test_clear_only_never_selects_ambiguous_pair_rule(cands):
    from monitoring_agent.exclusion import clear_only
    assert "rl_ueber_vl" in set(cands["rule_key"])                           # kommt in den Daten vor ...
    assert {rule for _, rule in clear_only(cands)} == {"nullwert", "zaehler_rueckgang"}   # ... wird aber nie automatisch gewählt
